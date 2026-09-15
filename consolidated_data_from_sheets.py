#!/usr/bin/env python3
"""
consolidate_data_from_sheets.py

Reads an .xlsx workbook and, for every sheet, produces a sparse "raw"
JSON file describing every non-empty cell: its coordinate, row/col,
cached value, formula text (if any), and merge range (if it's the
anchor of a merged region).

Design decisions (finalized):
  - Raw structure  : list of cell objects, ordered row-by-row (not a
                      cell-keyed dict) — easiest for an LLM/consumer to
                      scan sequentially.
  - Formula cells   : capture BOTH the formula text ("formula") and the
                      cached computed value ("value"). If the workbook
                      wasn't recalculated before saving, the cached value
                      may be missing; in that case "value" is null and
                      "value_stale_or_missing" is set to true.
  - Merged cells    : only the anchor (top-left) cell is emitted, tagged
                      with "merge_range" (e.g. "A1:A3"). The rest of the
                      merged region is skipped since it carries no data
                      of its own.

Usage:
    python consolidate_data_from_sheets.py path/to/workbook.xlsx [--outdir OUTDIR]

Output:
    OUTDIR/<sheetname>_raw.json   (one file per sheet)
"""

import argparse
import json
import os
import re
import sys

import openpyxl


def safe_filename(name: str) -> str:
    """Make a sheet name filesystem-safe for use in an output filename."""
    return re.sub(r"[^A-Za-z0-9_\-]+", "_", name).strip("_") or "sheet"


def extract_sheet(ws_values, ws_formulas):
    """
    ws_values:   worksheet loaded with data_only=True  (cached computed values)
    ws_formulas: worksheet loaded with data_only=False (raw formula text)

    Returns a list of cell dicts, ordered by row then column, containing
    only non-empty cells (and only merge-anchor cells within merged
    regions).
    """
    # Map anchor cell coord -> merge range string, and the full set of
    # coordinates that belong to some merged region (anchor + members).
    merge_map = {}
    merged_member_cells = set()
    for mrange in ws_values.merged_cells.ranges:
        anchor = mrange.coord.split(":")[0]
        merge_map[anchor] = str(mrange)
        for row in ws_values.iter_rows(
            min_row=mrange.min_row, max_row=mrange.max_row,
            min_col=mrange.min_col, max_col=mrange.max_col,
        ):
            for cell in row:
                merged_member_cells.add(cell.coordinate)

    cells = []
    for row_v, row_f in zip(ws_values.iter_rows(), ws_formulas.iter_rows()):
        for cell_v, cell_f in zip(row_v, row_f):
            coord = cell_v.coordinate

            # Skip non-anchor cells inside a merged region — they hold no
            # data of their own; the anchor cell carries the value.
            if coord in merged_member_cells and coord not in merge_map:
                continue

            is_formula = isinstance(cell_f.value, str) and cell_f.value.startswith("=")
            value = cell_v.value  # cached / computed value
            formula = cell_f.value if is_formula else None

            # Skip genuinely empty cells (no value, no formula).
            if value is None and formula is None:
                continue

            entry = {
                "cell": coord,
                "row": cell_v.row,
                "col": cell_v.column,
                "value": value,
            }
            if formula is not None:
                entry["formula"] = formula
                if value is None:
                    # Formula present but no cached value was saved in the
                    # file (e.g. workbook wasn't recalculated before save).
                    entry["value_stale_or_missing"] = True
            if coord in merge_map:
                entry["merge_range"] = merge_map[coord]

            cells.append(entry)

    return cells


def consolidate(xlsx_path: str, outdir: str):
    wb_values = openpyxl.load_workbook(xlsx_path, data_only=True)
    wb_formulas = openpyxl.load_workbook(xlsx_path, data_only=False)

    os.makedirs(outdir, exist_ok=True)
    written = []

    for sheetname in wb_values.sheetnames:
        ws_values = wb_values[sheetname]
        ws_formulas = wb_formulas[sheetname]

        cells = extract_sheet(ws_values, ws_formulas)

        out_path = os.path.join(outdir, f"{safe_filename(sheetname)}_raw.json")
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(
                {"sheet_name": sheetname, "cells": cells},
                f, indent=2, ensure_ascii=False, default=str,
            )
        written.append(out_path)
        print(f"  wrote {out_path}  ({len(cells)} cells)")

    return written


def main():
    parser = argparse.ArgumentParser(
        description="Extract per-sheet raw cell JSON from an xlsx workbook."
    )
    parser.add_argument("xlsx_path", help="Path to the input .xlsx workbook")
    parser.add_argument(
        "--outdir", default="./raw_sheets",
        help="Directory to write <sheetname>_raw.json files (default: ./raw_sheets)",
    )
    args = parser.parse_args()

    if not os.path.isfile(args.xlsx_path):
        sys.exit(f"File not found: {args.xlsx_path}")

    print(f"Reading {args.xlsx_path} ...")
    written = consolidate(args.xlsx_path, args.outdir)
    print(f"\nDone. {len(written)} sheet(s) written to {args.outdir}/")


if __name__ == "__main__":
    main()
